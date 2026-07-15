import pytest
from pathlib import Path
import fitz
from docx import Document

@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Hello from PDF\nThis is test content.", fontsize=12)
    path = tmp_path / "test.pdf"
    doc.save(path)
    doc.close()
    return path

@pytest.fixture
def sample_docx(tmp_path) -> Path:
    doc = Document()
    doc.add_heading("Test Heading", level=1)
    doc.add_paragraph("This is body text.")
    table = doc.add_table(rows=1, cols=2)
    row_cells = table.rows[0].cells
    row_cells[0].text = "Header 1"
    row_cells[1].text = "Header 2"
    path = tmp_path / "test.docx"
    doc.save(path)
    return path

@pytest.fixture
def sample_html(tmp_path) -> Path:
    html = """
    <html>
        <head><title>Test HTML</title></head>
        <body>
            <nav>Navigation links</nav>
            <main>
                <h1>Main Content</h1>
                <p>This is the main body.</p>
            </main>
            <footer>Footer notes</footer>
        </body>
    </html>
    """
    path = tmp_path / "test.html"
    path.write_text(html, encoding="utf-8")
    return path

@pytest.fixture
def sample_csv(tmp_path) -> Path:
    path = tmp_path / "test.csv"
    path.write_text("name,value\nfoo,1\nbar,2\n", encoding="utf-8")
    return path
