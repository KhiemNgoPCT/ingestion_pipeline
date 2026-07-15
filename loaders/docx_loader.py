from pathlib import Path
from docx import Document

def load_docx(path: Path):
    doc = Document(path)
    
    content = []
    
    # Process paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            content.append(para.text.strip())
            
    # Process tables
    table_count = len(doc.tables)
    for table in doc.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                content.append(" | ".join(row_data))
                
    raw_text = "\n".join(content)
    
    metadata = {
        "paragraph_count": len(doc.paragraphs),
        "table_count": table_count
    }
    
    return raw_text, "docx", "en", metadata
